import logging
import hashlib
import hmac
from functools import wraps
from io import BytesIO
from datetime import datetime, timezone

import openpyxl
from flask import Blueprint, abort, current_app, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_login import current_user, login_required, login_user, logout_user
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from werkzeug.security import check_password_hash, generate_password_hash

from extensions import db, login_manager
from activity import record_expense_created, record_expense_revised
from expense_history import (active_expense_records, append_expense_revision, claim_active_expense_record,
                             create_expense_record, iter_active_expenses, void_expense_record)
from helpers import (ZERO, compute_balances, compute_global_user_balances, compute_pairwise_user_obligations,
                     compute_settlements, compute_user_balances,
                     get_user_settlement_summary, parse_settlement_amount, save_splits, save_user_splits, validate_settlement_confirmation,
                     validate_settlement_request)
from models import (ActivityEvent, ConfirmedSettlement, Expense, ExpenseRecord, Group, GroupMembership, Member,
                    SettlementRequest, User)

main = Blueprint("main", __name__)
logger = logging.getLogger(__name__)


def group_member_or_error(group, member_id):
    member = db.session.get(Member, member_id)
    if member is None or member.group_id != group.id:
        raise ValueError("The selected payer must belong to this group.")
    return member


def group_users(group):
    """All linked members, including inactive users retained for history."""
    return [membership.user for membership in group.memberships]


def active_group_users(group):
    """Current members eligible for a new financial snapshot."""
    return (User.query.join(GroupMembership, GroupMembership.user_id == User.id)
            .filter(GroupMembership.group_id == group.id, User.is_active.is_(True))
            .order_by(GroupMembership.id).all())


def group_user_or_error(group, user_id):
    user = (User.query.join(GroupMembership, GroupMembership.user_id == User.id)
            .filter(GroupMembership.group_id == group.id, User.id == user_id, User.is_active.is_(True))
            .first())
    if user is None:
        raise ValueError("The selected payer must belong to this group and be active.")
    return user


def legacy_members_by_user(group, users):
    """Return explicit legacy compatibility rows for the group's global users.

    Missing rows are added with a known ``user_id``.  Existing name-only rows
    are never reinterpreted or changed.
    """
    members = {member.user_id: member for member in group.members if member.user_id is not None}
    for user in users:
        if user.id not in members:
            member = Member(name=user.name or user.login_id, group_id=group.id, user_id=user.id)
            db.session.add(member)
            db.session.flush()
            members[user.id] = member
    return members


def global_expense_context(group, expense=None):
    users = active_group_users(group)
    labels = {user.id: f"{user.name or user.login_id} — {user.login_id}" for user in users}
    if expense is None:
        return users, labels, [], None
    current_splits = {split.user_id: split.amount_owed for split in expense.splits if split.user_id is not None}
    return users, labels, list(current_splits), current_splits


def validate_expense_participant_payload(form, eligible_ids, label):
    """Reject forged participants even when their fields would be ignored."""
    shared_ids = {int(value) for value in form.getlist("shared_by")}
    if not shared_ids.issubset(eligible_ids):
        raise ValueError(f"Select valid {label} to share the expense; participants must be active.")
    for key in form:
        if key.startswith("custom_") and key != "custom_total":
            try:
                selected_id = int(key[len("custom_"):])
            except ValueError:
                raise ValueError("Select valid active group members.") from None
            if selected_id not in eligible_ids:
                raise ValueError(f"Select valid {label} to share the expense; participants must be active.")
        if key.startswith("share_") and key != "share_total":
            try:
                selected_id = int(key[len("share_"):])
            except ValueError:
                raise ValueError("Select valid active group members.") from None
            if selected_id not in eligible_ids:
                raise ValueError(f"Select valid {label} to share the expense; participants must be active.")


def active_legacy_members(group):
    """Mapped active members plus unmapped name-only historical members."""
    return (Member.query.outerjoin(User, Member.user_id == User.id)
            .filter(Member.group_id == group.id,
                    or_(Member.user_id.is_(None), User.is_active.is_(True))).all())


def expense_involves_inactive_user(expense):
    """An existing snapshot with an inactive identity is read-only."""
    user_ids = {expense.paid_by_user_id, expense.paid_by.user_id}
    for split in expense.splits:
        user_ids.update((split.user_id, split.member.user_id))
    user_ids.discard(None)
    return bool(user_ids and User.query.filter(User.id.in_(user_ids), User.is_active.is_(False)).first())


def build_expense_snapshot(group, form, global_identity):
    """Validate form data and build a new immutable Expense snapshot."""
    description = form.get("description", "").strip()
    expense_type = form.get("expense_type", "equal")
    if global_identity:
        users = active_group_users(group)
        validate_expense_participant_payload(form, {user.id for user in users}, "group users")
        payer = group_user_or_error(group, int(form.get("paid_by", "")))
        legacy_members = legacy_members_by_user(group, users)
        expense = Expense(description=description, paid_by_id=legacy_members[payer.id].id,
                          paid_by_user_id=payer.id, amount=ZERO, expense_type=expense_type, group_id=group.id)
        db.session.add(expense)
        db.session.flush()
        save_user_splits(expense, users, legacy_members, expense_type, form,
                         current_app.config["MAX_EXPENSE_AMOUNT"])
        return expense

    members = active_legacy_members(group)
    eligible_member_ids = {member.id for member in members}
    validate_expense_participant_payload(form, eligible_member_ids, "group members")
    paid_by_id = int(form.get("paid_by", ""))
    if paid_by_id not in eligible_member_ids:
        raise ValueError("The selected payer must belong to this group and be active.")
    expense = Expense(description=description, paid_by_id=paid_by_id, amount=ZERO,
                      expense_type=expense_type, group_id=group.id)
    db.session.add(expense)
    db.session.flush()
    save_splits(expense, members, expense_type, form, current_app.config["MAX_EXPENSE_AMOUNT"])
    return expense


def current_users_group(group_id):
    group = db.get_or_404(Group, group_id)
    is_member = GroupMembership.query.filter_by(group_id=group.id, user_id=current_user.id).first() is not None
    if group.owner_id != current_user.id and not is_member:
        abort(403)
    return group


def current_users_owned_group(group_id):
    group = current_users_group(group_id)
    if group.owner_id != current_user.id:
        abort(403)
    return group


def ensure_group_active(group):
    if group.is_archived:
        abort(403)
    return group


def require_expense_creator(record):
    """Authorize the current Stage 11 editing policy in one place."""
    if record.created_by_user_id != current_user.id:
        abort(403)
    return record


def claim_active_group(group):
    """Check archive state under the same writer lock as an authorized write."""
    claimed = Group.query.filter_by(id=group.id, archived_at=None).update(
        {Group.archived_at: Group.archived_at}, synchronize_session=False)
    if claimed != 1:
        db.session.rollback()
        abort(403)
    db.session.refresh(group)
    return group


def active_groups_for_user(user):
    return (Group.query.outerjoin(GroupMembership)
            .filter(or_(Group.owner_id == user.id, GroupMembership.user_id == user.id),
                    Group.archived_at.is_(None))
            .distinct().order_by(Group.id.desc()).all())


def admin_required(view):
    """Restrict the management control plane to the backend administrator."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        if current_user.is_authenticated:
            abort(403)
        if not admin_authenticated():
            return redirect(url_for("main.admin_login"))
        return view(*args, **kwargs)
    return wrapped


def admin_session_proof():
    username = current_app.config.get("ADMIN_USERNAME")
    password = current_app.config.get("ADMIN_PASSWORD")
    if not username or not password:
        return None
    credentials = f"{username}\0{password}".encode()
    return hmac.new(current_app.secret_key.encode(), credentials, hashlib.sha256).hexdigest()


def admin_authenticated():
    proof = admin_session_proof()
    return bool(proof and hmac.compare_digest(session.get("_admin_auth", ""), proof))


def valid_login_id(login_id):
    return len(login_id) == 10 and login_id.isascii() and login_id.isdigit()


def safe_next_url(value):
    if (not value or not value.startswith("/") or value.startswith("//")
            or "\\" in value or any(ord(char) < 32 or ord(char) == 127 for char in value)):
        return url_for("main.index")
    return value


@login_manager.user_loader
def load_user(user_id):
    try:
        user = db.session.get(User, int(user_id))
        return user if user is not None and user.is_active else None
    except (TypeError, ValueError):
        return None


@main.route("/login", methods=["GET", "POST"])
def login():
    if admin_authenticated():
        return redirect(url_for("main.admin_overview"))
    if current_user.is_authenticated:
        return redirect(url_for("main.index"))
    if request.method == "POST":
        login_id = request.form.get("login_id", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(login_id=login_id).first()
        if user and user.is_active and check_password_hash(user.password_hash, password):
            login_user(user)
            session.pop("_csrf_token", None)
            flash("You are signed in.", "success")
            return redirect(safe_next_url(request.args.get("next")))
        flash("Invalid ID or password.", "danger")
    return render_template("login.html")


@main.route("/register", methods=["GET", "POST"])
def register():
    abort(403)


@main.get("/users/search")
@login_required
def search_user():
    """Return one exact-ID match without exposing a browsable user directory."""
    login_id = request.args.get("login_id", "")
    if not valid_login_id(login_id):
        return jsonify(user=None)
    user = User.query.filter_by(login_id=login_id, is_active=True).first()
    return jsonify(user={"name": user.name or user.login_id, "login_id": user.login_id} if user else None)


@main.get("/admin")
@admin_required
def admin_overview():
    """User management and the existing read-only financial summary."""
    balances, unresolved = compute_global_user_balances()
    return render_template("admin.html", users=User.query.order_by(User.login_id).all(),
                           groups=Group.query.order_by(Group.id.desc()).all(), balances=balances,
                           unresolved_count=len(unresolved), zero=ZERO)


@main.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if current_user.is_authenticated:
        abort(403)
    if admin_authenticated():
        return redirect(url_for("main.admin_overview"))
    if request.method == "POST":
        username = current_app.config.get("ADMIN_USERNAME")
        configured_password = current_app.config.get("ADMIN_PASSWORD")
        supplied_username = request.form.get("username", "")
        supplied_password = request.form.get("password", "")
        if (username and configured_password and hmac.compare_digest(supplied_username, username)
                and hmac.compare_digest(supplied_password, configured_password)):
            session.clear()
            session["_admin_auth"] = admin_session_proof()
            flash("You are signed in as administrator.", "success")
            return redirect(url_for("main.admin_overview"))
        flash("Invalid credentials.", "danger")
    return render_template("admin_login.html")


@main.post("/admin/logout")
@admin_required
def admin_logout():
    session.clear()
    flash("You have been signed out.", "info")
    return redirect(url_for("main.admin_login"))


def new_user_validation_error(name, login_id, password, confirm_password):
    if not name:
        return "Name is required."
    if len(name) > 100:
        return "Name must be at most 100 characters."
    if not valid_login_id(login_id):
        return "User ID must contain exactly 10 digits."
    if User.query.filter_by(login_id=login_id).first() is not None:
        return "That User ID is already registered."
    if not password:
        return "Password is required."
    if password != confirm_password:
        return "Passwords do not match."
    return None


@main.post("/admin/users")
@admin_required
def admin_create_user():
    name = request.form.get("name", "").strip()
    login_id = request.form.get("login_id", "")
    password = request.form.get("password", "")
    error = new_user_validation_error(name, login_id, password, request.form.get("confirm_password", ""))
    if error:
        flash(error, "danger")
        return redirect(url_for("main.admin_overview"))
    try:
        db.session.add(User(name=name, login_id=login_id, password_hash=generate_password_hash(password),
                            is_active=True))
        db.session.commit()
        flash("User created.", "success")
    except IntegrityError:
        db.session.rollback()
        flash("That User ID is unavailable. Please try again.", "danger")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not create user")
        flash("We could not create that user. Please try again.", "danger")
    return redirect(url_for("main.admin_overview"))


@main.post("/admin/users/<int:user_id>/reset-password")
@admin_required
def admin_reset_password(user_id):
    user = db.get_or_404(User, user_id)
    password = request.form.get("password", "")
    if not password:
        flash("Password is required.", "danger")
    elif password != request.form.get("confirm_password", ""):
        flash("Passwords do not match.", "danger")
    else:
        try:
            user.password_hash = generate_password_hash(password)
            db.session.commit()
            flash("Password reset.", "success")
        except SQLAlchemyError:
            db.session.rollback()
            logger.exception("Could not reset user password")
            flash("We could not reset that password. Please try again.", "danger")
    return redirect(url_for("main.admin_overview"))


def deactivation_blocker(user_id):
    """Compose the existing ledger, direct-obligation, and request sources."""
    balances, _ = compute_global_user_balances()
    if balances.get(user_id, ZERO) != ZERO:
        return "This user has an outstanding global balance. Settle it first."
    obligations, _ = compute_pairwise_user_obligations()
    if any(user_id in pair and amount > ZERO for pair, amount in obligations.items()):
        return "This user has an outstanding direct obligation. Settle it first."
    pending = SettlementRequest.query.filter_by(status="PENDING").filter(
        or_(SettlementRequest.requester_user_id == user_id, SettlementRequest.recipient_user_id == user_id)
    ).first()
    if pending is not None:
        return "This user has a pending settlement request. Resolve it first."
    return None


@main.post("/admin/users/<int:user_id>/deactivate")
@admin_required
def admin_deactivate_user(user_id):
    user = db.get_or_404(User, user_id)
    if not user.is_active:
        flash("User is already deactivated.", "info")
        return redirect(url_for("main.admin_overview"))
    try:
        # Serialize the financial checks with SQLite ledger writers.
        claimed = User.query.filter_by(id=user_id, is_active=True).update(
            {User.is_active: User.is_active}, synchronize_session=False)
        if claimed != 1:
            db.session.rollback()
            flash("User is already deactivated.", "info")
            return redirect(url_for("main.admin_overview"))
        blocker = deactivation_blocker(user_id)
        if blocker:
            db.session.rollback()
            flash(blocker, "danger")
        else:
            user.is_active = False
            db.session.commit()
            flash("User deactivated. Financial history was preserved.", "success")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not deactivate user %s", user_id)
        flash("We could not deactivate that user. Please try again.", "danger")
    return redirect(url_for("main.admin_overview"))


@main.post("/admin/users/<int:user_id>/reactivate")
@admin_required
def admin_reactivate_user(user_id):
    user = db.get_or_404(User, user_id)
    if user.is_active:
        flash("User is already active.", "info")
        return redirect(url_for("main.admin_overview"))
    try:
        user.is_active = True
        db.session.commit()
        flash("User reactivated.", "success")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not reactivate user %s", user_id)
        flash("We could not reactivate that user. Please try again.", "danger")
    return redirect(url_for("main.admin_overview"))


@main.post("/settlement-requests")
@login_required
def create_settlement_request():
    try:
        recipient_user_id = int(request.form.get("recipient_user_id", ""))
        recipient = db.session.get(User, recipient_user_id)
        if recipient is None:
            raise ValueError("Recipient user was not found.")
        amount = parse_settlement_amount(request.form.get("amount", ""))
        # Serialize pending-limit validation and insertion on SQLite. This is
        # a no-op claim, not a change to the user's identity or permissions.
        User.query.filter_by(id=current_user.id).update({User.id: User.id}, synchronize_session=False)
        validate_settlement_request(current_user.id, recipient.id, amount)
        db.session.add(SettlementRequest(requester_user_id=current_user.id, recipient_user_id=recipient.id,
                                         amount=amount, status="PENDING"))
        db.session.commit()
        flash("Settlement request created.", "success")
    except (ValueError, TypeError) as error:
        db.session.rollback()
        flash(str(error), "danger")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not create settlement request")
        flash("We could not create that settlement request. Please try again.", "danger")
    return redirect(url_for("main.index"))


def resolve_settlement_request(request_id, target_status):
    settlement_request = db.get_or_404(SettlementRequest, request_id)
    if settlement_request.recipient_user_id != current_user.id:
        abort(403)
    if settlement_request.status != "PENDING":
        flash("Only pending settlement requests can be resolved.", "danger")
        return redirect(url_for("main.index"))
    try:
        # A conditional no-op UPDATE claims the pending row before reading the
        # ledger. On SQLite this also serializes writers, including confirmations
        # of different requests. Keep the claim and both writes in one transaction.
        claimed = SettlementRequest.query.filter_by(
            id=request_id, recipient_user_id=current_user.id, status="PENDING"
        ).update({SettlementRequest.status: "PENDING"}, synchronize_session=False)
        if claimed != 1:
            raise ValueError("Only pending settlement requests can be resolved.")
        db.session.refresh(settlement_request)
        resolved_at = datetime.now(timezone.utc)
        if target_status == "CONFIRMED":
            validate_settlement_confirmation(settlement_request)
            db.session.add(ConfirmedSettlement(
                settlement_request_id=settlement_request.id,
                payer_user_id=settlement_request.requester_user_id,
                recipient_user_id=settlement_request.recipient_user_id,
                amount=settlement_request.amount,
                confirmed_at=resolved_at,
            ))
        settlement_request.status = target_status
        settlement_request.resolved_at = resolved_at
        db.session.commit()
        flash(f"Settlement request {target_status.lower()}.", "success")
    except ValueError as error:
        db.session.rollback()
        flash(str(error), "danger")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not resolve settlement request %s", request_id)
        flash("We could not resolve that settlement request. Please try again.", "danger")
    return redirect(url_for("main.index"))


@main.post("/settlement-requests/<int:request_id>/confirm")
@login_required
def confirm_settlement_request(request_id):
    return resolve_settlement_request(request_id, "CONFIRMED")


@main.post("/settlement-requests/<int:request_id>/reject")
@login_required
def reject_settlement_request(request_id):
    return resolve_settlement_request(request_id, "REJECTED")


@main.post("/logout")
@login_required
def logout():
    logout_user()
    session.pop("_csrf_token", None)
    flash("You have been signed out.", "info")
    return redirect(url_for("main.login"))


@main.route("/")
@login_required
def index():
    groups = active_groups_for_user(current_user)
    eligible_users = (User.query
                      .filter(User.id != current_user.id, User.is_active.is_(True))
                      .order_by(User.name, User.login_id)
                      .all())
    active_counts = {group.id: len(active_expense_records(group)) for group in groups}
    settlement_requests = (SettlementRequest.query
                           .filter(or_(SettlementRequest.requester_user_id == current_user.id,
                                       SettlementRequest.recipient_user_id == current_user.id))
                           .order_by(SettlementRequest.created_at.desc(), SettlementRequest.id.desc()).all())
    outgoing_summary, incoming_summary = get_user_settlement_summary(current_user.id)
    partner_ids = {row["user_id"] for row in outgoing_summary + incoming_summary}
    partners = {user.id: user for user in User.query.filter(User.id.in_(partner_ids)).all()} if partner_ids else {}
    outgoing_obligations = [{**row, "user": partners[row["user_id"]]} for row in outgoing_summary]
    incoming_obligations = [{**row, "user": partners[row["user_id"]]} for row in incoming_summary]
    outgoing_obligations.sort(key=lambda row: (row["user"].name or row["user"].login_id, row["user"].id))
    incoming_obligations.sort(key=lambda row: (row["user"].name or row["user"].login_id, row["user"].id))
    return render_template("index.html", groups=groups, active_counts=active_counts,
                           eligible_users=eligible_users,
                           settlement_requests=settlement_requests,
                           outgoing_obligations=outgoing_obligations,
                           incoming_obligations=incoming_obligations, zero=ZERO)


@main.get("/activity")
@login_required
def activity_history():
    events = (ActivityEvent.query
              .join(Group, ActivityEvent.group_id == Group.id)
              .outerjoin(GroupMembership, GroupMembership.group_id == ActivityEvent.group_id)
              .filter(or_(Group.owner_id == current_user.id,
                          GroupMembership.user_id == current_user.id))
              .distinct()
              .order_by(ActivityEvent.created_at.desc(), ActivityEvent.id.desc())
              .all())
    return render_template("activity.html", events=events)


@main.route("/groups", methods=["POST"])
@login_required
def create_group():
    name = request.form.get("name", "").strip()
    login_ids = list(dict.fromkeys(value for value in request.form.getlist("user_ids") if value))
    if not name:
        flash("Enter a group name.", "danger")
        return redirect(url_for("main.index"))
    if len(name) > 100:
        flash("Group name must be at most 100 characters.", "danger")
        return redirect(url_for("main.index"))
    if any(not valid_login_id(login_id) for login_id in login_ids):
        flash("Selected User IDs must contain exactly 10 digits.", "danger")
        return redirect(url_for("main.index"))

    selected_users = User.query.filter(User.login_id.in_(login_ids), User.is_active.is_(True)).all() if login_ids else []
    if len(selected_users) != len(login_ids):
        flash("One or more selected User IDs are no longer registered.", "danger")
        return redirect(url_for("main.index"))
    users_by_id = {user.login_id: user for user in selected_users}
    users_by_id[current_user.login_id] = current_user

    try:
        group = Group(name=name, owner_id=current_user.id)
        db.session.add(group)
        db.session.flush()
        selected_users = list(users_by_id.values())
        db.session.add_all(GroupMembership(group_id=group.id, user_id=user.id) for user in selected_users)
        # Required compatibility rows for the immutable Expense/ExpenseSplit snapshots.
        db.session.add_all(Member(name=user.name or user.login_id, group_id=group.id, user_id=user.id) for user in selected_users)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not create group")
        flash("We could not create that group. Please try again.", "danger")
        return redirect(url_for("main.index"))
    flash(f'Group "{name}" created!', "success")
    return redirect(url_for("main.view_group", group_id=group.id))


@main.route("/groups/<int:group_id>")
@login_required
def view_group(group_id):
    group = current_users_group(group_id)
    expense_records = active_expense_records(group)
    balances = compute_balances(group)
    show_group_suggested_payments = current_app.config["SHOW_GROUP_SUGGESTED_PAYMENTS"]
    settlements = compute_settlements(balances) if show_group_suggested_payments else None
    user_balances = compute_user_balances(group)
    has_unmapped_legacy_expenses = any(
        expense.paid_by_user_id is None or any(split.user_id is None for split in expense.splits)
        for expense in iter_active_expenses(group)
    )
    return render_template("group.html", group=group, balances=balances,
                           settlements=settlements,
                           show_group_suggested_payments=show_group_suggested_payments,
                           member_map={member.id: member.name for member in group.members}, zero=ZERO,
                           is_owner=(group.owner_id == current_user.id), user_balances=user_balances,
                           has_unmapped_legacy_expenses=has_unmapped_legacy_expenses,
                           expense_records=expense_records, is_archived=group.is_archived)


@main.route("/groups/<int:group_id>/expenses/new", methods=["GET", "POST"])
@login_required
def add_expense(group_id):
    group = ensure_group_active(current_users_group(group_id))
    users, participant_labels, _, _ = global_expense_context(group)
    if request.method == "POST":
        try:
            claim_active_group(group)
            expense = build_expense_snapshot(group, request.form, global_identity=True)
            record = create_expense_record(expense, current_user.id)
            record_expense_created(record, expense, current_user.id)
            db.session.commit()
            flash("Expense added!", "success")
        except (ValueError, TypeError) as error:
            db.session.rollback()
            flash(str(error), "danger")
        except SQLAlchemyError:
            db.session.rollback()
            logger.exception("Could not add expense to group %s", group_id)
            flash("We could not save that expense. Please try again.", "danger")
        return redirect(url_for("main.view_group", group_id=group.id))
    return render_template("add_expense.html", group=group, participants=users,
                           participant_labels=participant_labels)


@main.route("/groups/<int:group_id>/expenses/<int:expense_id>/edit", methods=["GET", "POST"])
@login_required
def edit_expense(group_id, expense_id):
    group = ensure_group_active(current_users_group(group_id))
    record = db.get_or_404(ExpenseRecord, expense_id)
    if record.group_id != group.id:
        abort(404)
    require_expense_creator(record)
    if record.current_expense_id is None:
        flash("Voided expenses cannot be edited.", "danger")
        return redirect(url_for("main.view_group", group_id=group.id))
    expense = record.current_expense
    if expense_involves_inactive_user(expense):
        flash("This expense involves a deactivated user and cannot be revised while they are inactive.", "danger")
        return redirect(url_for("main.view_group", group_id=group.id))
    is_global_expense = expense.paid_by_user_id is not None
    if request.method == "POST":
        try:
            claim_active_group(group)
            record = claim_active_expense_record(record.id)
            before = record.current_expense
            if expense_involves_inactive_user(before):
                raise ValueError("This expense involves a deactivated user and cannot be revised while they are inactive.")
            snapshot = build_expense_snapshot(group, request.form,
                                              global_identity=(record.current_expense.paid_by_user_id is not None))
            changed = append_expense_revision(record, snapshot, current_user.id)
            if changed:
                record_expense_revised(record, before, snapshot, current_user.id)
            db.session.commit()
            flash("Expense updated!" if changed else "No expense changes were detected.",
                  "success" if changed else "info")
        except (ValueError, TypeError) as error:
            db.session.rollback()
            flash(str(error), "danger")
        except SQLAlchemyError:
            db.session.rollback()
            logger.exception("Could not update expense %s", expense_id)
            flash("We could not save those changes. Please try again.", "danger")
        return redirect(url_for("main.view_group", group_id=group.id))
    if is_global_expense:
        participants, participant_labels, shared_ids, current_splits = global_expense_context(group, expense)
        selected_payer_id = expense.paid_by_user_id
        current_shares = {split.user_id: split.share_weight for split in expense.splits
                          if split.user_id is not None and split.share_weight is not None}
    else:
        participants = active_legacy_members(group)
        participant_labels = {member.id: member.name for member in participants}
        current_splits = {split.member_id: split.amount_owed for split in expense.splits}
        shared_ids = list(current_splits)
        selected_payer_id = expense.paid_by_id
        current_shares = {split.member_id: split.share_weight for split in expense.splits
                          if split.share_weight is not None}
    return render_template("edit_expense.html", group=group, record=record, expense=expense, participants=participants,
                           participant_labels=participant_labels, current_splits=current_splits,
                           current_shares=current_shares, shared_ids=shared_ids, selected_payer_id=selected_payer_id)


@main.post("/groups/<int:group_id>/expenses/<int:expense_id>/delete")
@login_required
def delete_expense(group_id, expense_id):
    group = ensure_group_active(current_users_owned_group(group_id))
    record = db.get_or_404(ExpenseRecord, expense_id)
    if record.group_id != group.id:
        abort(404)
    try:
        claim_active_group(group)
        void_expense_record(record.id, current_user.id)
        db.session.commit()
        flash("Expense voided. Its history has been preserved.", "info")
    except ValueError as error:
        db.session.rollback()
        flash(str(error), "danger")
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not void expense record %s", expense_id)
        flash("We could not void that expense. Please try again.", "danger")
    return redirect(url_for("main.view_group", group_id=group_id))


@main.get("/groups/<int:group_id>/expenses/<int:expense_id>/history")
@login_required
def expense_history(group_id, expense_id):
    group = current_users_group(group_id)
    record = db.get_or_404(ExpenseRecord, expense_id)
    if record.group_id != group.id:
        abort(404)
    return render_template("expense_history.html", group=group, record=record)


@main.post("/groups/<int:group_id>/delete")
@login_required
def delete_group(group_id):
    group = current_users_owned_group(group_id)
    if group.is_archived:
        flash("This group is already archived.", "info")
        return redirect(url_for("main.index"))
    try:
        changed = Group.query.filter_by(id=group.id, owner_id=current_user.id, archived_at=None).update(
            {Group.archived_at: datetime.now(timezone.utc), Group.archived_by_user_id: current_user.id},
            synchronize_session=False)
        db.session.commit()
        if not changed:
            flash("This group is already archived.", "info")
            return redirect(url_for("main.index"))
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception("Could not archive group %s", group_id)
        flash("We could not archive that group. Please try again.", "danger")
        return redirect(url_for("main.index"))
    flash(f'Group "{group.name}" archived. Its financial history is preserved.', "info")
    return redirect(url_for("main.index"))


@main.route("/groups/<int:group_id>/export")
@login_required
def export_excel(group_id):
    group = current_users_group(group_id)
    balances = compute_balances(group)
    settlements = compute_settlements(balances)
    member_map = {member.id: member.name for member in group.members}
    workbook = openpyxl.Workbook()
    expenses_sheet = workbook.active
    expenses_sheet.title = "Expenses Log"
    expenses_sheet.append(["#", "Date", "Description", "Paid By", "Amount (₹)", "Split Among", "Type"])
    for index, expense in enumerate(iter_active_expenses(group), 1):
        expenses_sheet.append([index, expense.created_at, expense.description or "", expense.paid_by.name,
                               expense.amount, ", ".join(f"{member_map[split.member_id]} (₹{split.amount_owed:.2f})" for split in expense.splits),
                               expense.expense_type.capitalize()])
    balances_sheet = workbook.create_sheet("Final Balances")
    balances_sheet.append(["Member", "Balance (₹)", "Status"])
    for member in group.members:
        balance = balances.get(member.id, ZERO)
        balances_sheet.append([member.name, abs(balance), "Settled" if balance == ZERO else "Should Receive" if balance > ZERO else "Should Pay"])
    payments_sheet = workbook.create_sheet("Suggested Payments")
    payments_sheet.append(["From (Debtor)", "To (Creditor)", "Amount (₹)"])
    for debtor_id, creditor_id, amount in settlements:
        payments_sheet.append([member_map[debtor_id], member_map[creditor_id], amount])
    # This export has no intentional formulas. Preserve user text as text even
    # when it starts with '=' (openpyxl otherwise treats it as an Excel formula).
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=f"{group.name.replace(' ', '_')}_expenses.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
