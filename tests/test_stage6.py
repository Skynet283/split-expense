import unittest
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier
from unittest.mock import patch

import openpyxl
from sqlalchemy.exc import IntegrityError
from werkzeug.security import generate_password_hash

from app import create_app, initialize_database
from config import Config
from extensions import db
from expense_history import iter_active_expenses
from helpers import compute_global_user_balances, compute_pairwise_user_obligations
from models import (ConfirmedSettlement, Expense, ExpenseRecord, ExpenseRevision, ExpenseSplit, Group,
                    GroupMembership, Member, SettlementRequest, User)
from test_factory import create_user


class TestConfig(Config):
    TESTING = True
    CSRF_ENABLED = False
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    INITIAL_ADMIN_ID = "9999999999"
    INITIAL_ADMIN_PASSWORD = "test-admin-password"
    INITIAL_ADMIN_NAME = "Test Administrator"
    ADMIN_USERNAME = "test-admin"
    ADMIN_PASSWORD_HASH = generate_password_hash("test-admin-password")


class Stage6Tests(unittest.TestCase):
    def setUp(self):
        self.app = create_app(TestConfig)
        with self.app.app_context():
            db.drop_all()
            db.create_all()
        self.client = self.app.test_client()

    def register(self, name, login_id):
        create_user(self.app, name, login_id)

    def login(self, login_id):
        self.client.post("/login", data={"login_id": login_id, "password": "secret"})

    def seed(self, amount="500.00"):
        self.register("Asha", "0123456789")
        self.register("Bala", "1234567890")
        self.register("Chitra", "2345678901")
        self.login("0123456789")
        self.client.post("/groups", data={"name": "Trip", "user_ids": ["1234567890", "2345678901"]})
        with self.app.app_context():
            group = Group.query.filter_by(name="Trip").one()
            users = {user.login_id: user.id for user in User.query.all()}
            group_id = group.id
        self.client.post(f"/groups/{group_id}/expenses/new", data={
            "description": "Dinner", "paid_by": str(users["0123456789"]), "expense_type": "custom",
            "custom_total": amount, f"custom_{users['1234567890']}": amount,
        })
        with self.app.app_context():
            return group_id, ExpenseRecord.query.one().id, users

    def edit(self, group_id, record_id, users, *, description="Dinner", payer="0123456789",
             shares=None, total="500.00"):
        shares = shares or {"1234567890": total}
        data = {"description": description, "paid_by": str(users[payer]), "expense_type": "custom",
                "custom_total": total}
        data.update({f"custom_{users[login_id]}": amount for login_id, amount in shares.items()})
        return self.client.post(f"/groups/{group_id}/expenses/{record_id}/edit", data=data,
                                follow_redirects=True)

    def test_existing_expense_backfill_is_exact_and_idempotent(self):
        with self.app.app_context():
            owner = User(name="Legacy Owner", login_id="0123456789", password_hash="hash")
            db.session.add(owner)
            db.session.flush()
            group = Group(name="Legacy", owner_id=owner.id)
            db.session.add(group)
            db.session.flush()
            member = Member(name="Name Only", group_id=group.id)
            db.session.add(member)
            db.session.flush()
            expense = Expense(description="Old", paid_by_id=member.id, amount=Decimal("12.34"),
                              group_id=group.id, expense_type="custom")
            db.session.add(expense)
            db.session.flush()
            db.session.add(ExpenseSplit(expense_id=expense.id, member_id=member.id,
                                        amount_owed=Decimal("12.34")))
            db.session.commit()
            expense_id, group_id, member_id = expense.id, group.id, member.id

        initialize_database(self.app)
        initialize_database(self.app)
        with self.app.app_context():
            record = ExpenseRecord.query.one()
            revision = ExpenseRevision.query.one()
            snapshot = db.session.get(Expense, expense_id)
            self.assertEqual(record.id, expense_id)
            self.assertEqual(record.group_id, group_id)
            self.assertEqual(record.current_expense_id, expense_id)
            self.assertEqual(revision.version_number, 1)
            self.assertEqual(snapshot.description, "Old")
            self.assertEqual(snapshot.amount, Decimal("12.34"))
            self.assertEqual(snapshot.paid_by_id, member_id)
            self.assertIsNone(snapshot.paid_by_user_id)
            self.assertIsNone(snapshot.splits[0].user_id)
            self.assertEqual(ExpenseRecord.query.count(), 1)
            self.assertEqual(ExpenseRevision.query.count(), 1)

    def test_edits_append_complete_versions_and_identical_edit_is_noop(self):
        group_id, record_id, users = self.seed()
        with self.app.app_context():
            original = ExpenseRecord.query.one().current_expense
            original_id = original.id
            original_values = (original.description, original.paid_by_user_id, original.amount,
                               [(split.user_id, split.amount_owed) for split in original.splits])

        self.edit(group_id, record_id, users, description="Corrected", payer="2345678901",
                  shares={"0123456789": "100.00", "1234567890": "600.00"}, total="700.00")
        with self.app.app_context():
            record = db.session.get(ExpenseRecord, record_id)
            current = record.current_expense
            self.assertNotEqual(current.id, original_id)
            self.assertEqual(current.description, "Corrected")
            self.assertEqual(current.paid_by_user_id, users["2345678901"])
            self.assertEqual(current.amount, Decimal("700.00"))
            self.assertEqual({(s.user_id, s.amount_owed) for s in current.splits},
                             {(users["0123456789"], Decimal("100.00")),
                              (users["1234567890"], Decimal("600.00"))})
            old = db.session.get(Expense, original_id)
            self.assertEqual((old.description, old.paid_by_user_id, old.amount,
                              [(s.user_id, s.amount_owed) for s in old.splits]), original_values)
            self.assertEqual([r.version_number for r in record.revisions], [1, 2])

        self.edit(group_id, record_id, users, description="Corrected", payer="2345678901",
                  shares={"0123456789": "100.00", "1234567890": "600.00"}, total="700.00")
        with self.app.app_context():
            self.assertEqual(ExpenseRevision.query.count(), 2)
            self.assertEqual(Expense.query.count(), 2)

        self.edit(group_id, record_id, users, description="Receipt clarified", payer="2345678901",
                  shares={"0123456789": "100.00", "1234567890": "600.00"}, total="700.00")
        with self.app.app_context():
            self.assertEqual(ExpenseRevision.query.count(), 3)
            self.assertEqual(ExpenseRecord.query.one().current_expense.description, "Receipt clarified")

    def test_only_active_snapshot_drives_global_and_pairwise_accounting(self):
        group_id, record_id, users = self.seed()
        self.edit(group_id, record_id, users, total="700.00", shares={"1234567890": "700.00"})
        with self.app.app_context():
            balances, unresolved = compute_global_user_balances()
            obligations, pairwise_unresolved = compute_pairwise_user_obligations()
            self.assertEqual(balances[users["0123456789"]], Decimal("700.00"))
            self.assertEqual(balances[users["1234567890"]], Decimal("-700.00"))
            self.assertEqual(sum(balances.values()), Decimal("0.00"))
            self.assertEqual(obligations, {(users["1234567890"], users["0123456789"]): Decimal("700.00")})
            self.assertEqual(unresolved, pairwise_unresolved)
            self.assertEqual(len(list(iter_active_expenses())), 1)
            self.assertEqual(Expense.query.count(), 2)

    def test_edit_after_settlement_preserves_payment_and_can_reverse_obligation(self):
        group_id, record_id, users = self.seed()
        self.client.post("/logout")
        self.login("1234567890")
        self.client.post("/settlement-requests", data={"recipient_user_id": users["0123456789"],
                                                        "amount": "500.00"})
        with self.app.app_context():
            request_id = SettlementRequest.query.one().id
        self.client.post("/logout")
        self.login("0123456789")
        self.client.post(f"/settlement-requests/{request_id}/confirm")
        self.edit(group_id, record_id, users, total="700.00", shares={"1234567890": "700.00"})
        with self.app.app_context():
            payment = ConfirmedSettlement.query.one()
            self.assertEqual(payment.amount, Decimal("500.00"))
            self.assertEqual(compute_pairwise_user_obligations()[0],
                             {(users["1234567890"], users["0123456789"]): Decimal("200.00")})
        self.edit(group_id, record_id, users, total="300.00", shares={"1234567890": "300.00"})
        with self.app.app_context():
            self.assertEqual(ConfirmedSettlement.query.one().amount, Decimal("500.00"))
            self.assertEqual(compute_pairwise_user_obligations()[0],
                             {(users["0123456789"], users["1234567890"]): Decimal("200.00")})

    def test_void_preserves_history_and_payment_and_rejects_future_changes(self):
        group_id, record_id, users = self.seed()
        self.client.post("/logout")
        self.login("1234567890")
        self.client.post("/settlement-requests", data={"recipient_user_id": users["0123456789"],
                                                        "amount": "500.00"})
        with self.app.app_context():
            request_id = SettlementRequest.query.one().id
        self.client.post("/logout")
        self.login("0123456789")
        self.client.post(f"/settlement-requests/{request_id}/confirm")
        response = self.client.post(f"/groups/{group_id}/expenses/{record_id}/delete", follow_redirects=True)
        self.assertIn(b"history has been preserved", response.data)
        second = self.client.post(f"/groups/{group_id}/expenses/{record_id}/delete", follow_redirects=True)
        self.assertIn(b"Only an active expense", second.data)
        edit = self.edit(group_id, record_id, users, total="600.00", shares={"1234567890": "600.00"})
        self.assertIn(b"Voided expenses cannot be edited", edit.data)
        with self.app.app_context():
            record = db.session.get(ExpenseRecord, record_id)
            self.assertIsNone(record.current_expense_id)
            self.assertIsNotNone(record.voided_at)
            self.assertEqual(record.voided_by_user_id, users["0123456789"])
            self.assertEqual(Expense.query.count(), 1)
            self.assertEqual(ExpenseRevision.query.count(), 1)
            self.assertEqual(ConfirmedSettlement.query.one().amount, Decimal("500.00"))
            self.assertEqual(compute_pairwise_user_obligations()[0],
                             {(users["0123456789"], users["1234567890"]): Decimal("500.00")})

    def test_history_authorization_and_group_archive_policy(self):
        group_id, record_id, users = self.seed()
        history = self.client.get(f"/groups/{group_id}/expenses/{record_id}/history")
        self.assertIn(b"Version 1", history.data)
        self.client.post("/logout")
        self.login("1234567890")
        self.assertEqual(self.client.get(f"/groups/{group_id}/expenses/{record_id}/history").status_code, 200)
        self.assertEqual(self.client.get(f"/groups/{group_id}/expenses/{record_id}/edit").status_code, 403)
        self.assertEqual(self.client.post(f"/groups/{group_id}/expenses/{record_id}/delete").status_code, 403)
        self.client.post("/logout")
        self.login("0123456789")
        archived = self.client.post(f"/groups/{group_id}/delete", follow_redirects=True)
        self.assertIn(b"archived", archived.data)
        with self.app.app_context():
            self.assertIsNotNone(db.session.get(Group, group_id).archived_at)
            empty = Group(name="Empty", owner_id=users["0123456789"])
            db.session.add(empty)
            db.session.commit()
            empty_id = empty.id
        self.client.post(f"/groups/{empty_id}/delete")
        with self.app.app_context():
            self.assertIsNotNone(db.session.get(Group, empty_id))
            self.assertIsNotNone(db.session.get(Group, empty_id).archived_at)

    def test_current_display_and_export_exclude_old_and_voided_versions(self):
        group_id, record_id, users = self.seed()
        self.edit(group_id, record_id, users, description="Current only", total="700.00",
                  shares={"1234567890": "700.00"})
        page = self.client.get(f"/groups/{group_id}")
        self.assertIn(b"Current only", page.data)
        self.assertNotIn(b">Dinner<", page.data)
        workbook = openpyxl.load_workbook(BytesIO(self.client.get(f"/groups/{group_id}/export").data))
        rows = list(workbook["Expenses Log"].iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], "Current only")
        self.client.post(f"/groups/{group_id}/expenses/{record_id}/delete")
        page = self.client.get(f"/groups/{group_id}")
        self.assertNotIn(b"Current only", page.data)
        workbook = openpyxl.load_workbook(BytesIO(self.client.get(f"/groups/{group_id}/export").data))
        self.assertEqual(len(list(workbook["Expenses Log"].iter_rows(values_only=True))), 1)

    def test_revision_database_constraints(self):
        group_id, record_id, users = self.seed()
        with self.app.app_context():
            record = db.session.get(ExpenseRecord, record_id)
            current_id = record.current_expense_id
            for version, expense_id in ((0, current_id), (1, current_id)):
                db.session.add(ExpenseRevision(expense_record_id=record_id, expense_id=expense_id,
                                               version_number=version))
                with self.assertRaises(IntegrityError):
                    db.session.commit()
                db.session.rollback()
            self.assertEqual(ExpenseRevision.query.count(), 1)

    def test_concurrent_edits_form_a_single_ordered_history(self):
        with TemporaryDirectory() as directory:
            class FileConfig(TestConfig):
                SQLALCHEMY_DATABASE_URI = "sqlite:///" + (Path(directory) / "history.db").as_posix()

            self.app = create_app(FileConfig)
            self.client = self.app.test_client()
            try:
                group_id, record_id, users = self.seed()
                clients = [self.app.test_client(), self.app.test_client()]
                for client in clients:
                    client.post("/login", data={"login_id": "0123456789", "password": "secret"})
                barrier = Barrier(2)
                original_get = db.get_or_404

                def simultaneous_read(model, identifier, **kwargs):
                    row = original_get(model, identifier, **kwargs)
                    if model is ExpenseRecord:
                        barrier.wait(timeout=10)
                    return row

                forms = [
                    {"description": "Revision A", "paid_by": users["0123456789"], "expense_type": "custom",
                     "custom_total": "600.00", f"custom_{users['1234567890']}": "600.00"},
                    {"description": "Revision B", "paid_by": users["0123456789"], "expense_type": "custom",
                     "custom_total": "700.00", f"custom_{users['1234567890']}": "700.00"},
                ]
                with patch.object(db, "get_or_404", side_effect=simultaneous_read):
                    with ThreadPoolExecutor(max_workers=2) as executor:
                        futures = [executor.submit(client.post,
                                                   f"/groups/{group_id}/expenses/{record_id}/edit", data=form)
                                   for client, form in zip(clients, forms)]
                        responses = [future.result(timeout=15) for future in futures]
                self.assertTrue(all(response.status_code == 302 for response in responses))
                with self.app.app_context():
                    record = db.session.get(ExpenseRecord, record_id)
                    self.assertEqual([revision.version_number for revision in record.revisions], [1, 2, 3])
                    self.assertIn(record.current_expense.description, {"Revision A", "Revision B"})
                    self.assertEqual(len(list(iter_active_expenses())), 1)
            finally:
                with self.app.app_context():
                    db.session.remove()
                    db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
